from io import BytesIO
from pathlib import Path

from app.bootstrap.api import create_app
from app.core.settings import ApiSettings
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _xlsx_bytes(text: str) -> bytes:
    workbook = Workbook()
    workbook.active.append([text])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _settings(tmp_path: Path) -> ApiSettings:
    return ApiSettings(
        repository_root=tmp_path,
        backend_root=tmp_path / "backend",
        data_root=tmp_path / "data",
        cors_origins=(),
        run_workers=1,
        database_url=f"sqlite+pysqlite:///{tmp_path / 'api.db'}",
        storage_root=tmp_path / "storage",
        upload_max_files=20,
        upload_max_bytes=100_000_000,
    )


def _valid_files() -> dict[str, bytes]:
    return {
        "AWP/scope.xlsx": _xlsx_bytes("Approved audit scope"),
        "APM/risk.docx": _docx_bytes("Planning risk context"),
        "Process Understanding/evidence.docx": _docx_bytes(
            "Observed process evidence"
        ),
        "Process SOP/criteria.docx": _docx_bytes(
            "Approved control criteria"
        ),
    }


def _create_session(
    client: TestClient,
    payloads: dict[str, bytes],
):
    response = client.post(
        "/api/v1/upload-sessions",
        json={
            "files": [
                {
                    "relative_path": path,
                    "size_bytes": len(content),
                    "content_type": _CONTENT_TYPE,
                }
                for path, content in payloads.items()
            ]
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _upload_all(
    client: TestClient,
    session: dict,
    payloads: dict[str, bytes],
) -> None:
    for file in session["files"]:
        response = client.put(
            file["upload_url"],
            content=payloads[file["relative_path"]],
            headers={"Content-Type": _CONTENT_TYPE},
        )
        assert response.status_code == 200, response.text
        assert response.json()["upload_status"] == "UPLOADED"


def test_local_upload_session_validates_and_creates_v01(
    tmp_path: Path,
) -> None:
    app = create_app(settings=_settings(tmp_path))
    payloads = _valid_files()

    with TestClient(app) as client:
        session = _create_session(client, payloads)
        assert session["state"] == "UPLOADING"
        assert session["allowed_actions"] == [
            "DISCARD",
            "UPLOAD_FILES",
            "VALIDATE",
        ]

        _upload_all(client, session, payloads)
        validated = client.post(
            f"/api/v1/upload-sessions/{session['session_id']}/validate"
        )
        assert validated.status_code == 200, validated.text
        report = validated.json()["validation_report"]
        assert validated.json()["state"] == "READY_TO_CREATE"
        assert report["valid"] is True
        assert report["role_summary"] == {
            "SCOPE": 1,
            "RISK_CONTEXT": 1,
            "EVIDENCE": 1,
            "CRITERIA": 1,
        }

        promoted = client.post(
            f"/api/v1/upload-sessions/{session['session_id']}/projects",
            json={"name": "FY2026 Local Intake"},
        )
        assert promoted.status_code == 201, promoted.text
        project = promoted.json()
        assert project["state"] == "READY_FOR_DISCOVERY"
        assert project["version"]["label"] == "v0.1"

        versions = client.get(
            f"/api/v1/projects/{project['project_id']}/versions"
        )
        assert versions.status_code == 200, versions.text
        assert [item["label"] for item in versions.json()] == ["v0.1"]

        cannot_discard = client.delete(
            f"/api/v1/upload-sessions/{session['session_id']}"
        )
        assert cannot_discard.status_code == 409
        assert cannot_discard.json()["error"]["code"] == "INVALID_STATE"

    source_root = (
        tmp_path
        / "storage"
        / "uat-intake"
        / "projects"
        / project["project_id"]
        / "source"
    )
    assert {
        path.relative_to(source_root).as_posix()
        for path in source_root.rglob("*")
        if path.is_file()
    } == set(payloads)
    assert not (
        tmp_path
        / "storage"
        / "uat-intake"
        / "staging"
        / session["session_id"]
    ).exists()


def test_validation_reports_missing_files_then_allows_reupload(
    tmp_path: Path,
) -> None:
    app = create_app(settings=_settings(tmp_path))
    payloads = _valid_files()

    with TestClient(app) as client:
        session = _create_session(client, payloads)
        first = session["files"][0]
        uploaded = client.put(
            first["upload_url"],
            content=payloads[first["relative_path"]],
            headers={"Content-Type": _CONTENT_TYPE},
        )
        assert uploaded.status_code == 200

        invalid = client.post(
            f"/api/v1/upload-sessions/{session['session_id']}/validate"
        )
        assert invalid.status_code == 200
        assert invalid.json()["state"] == "INVALID"
        assert invalid.json()["validation_report"]["valid"] is False
        assert any(
            error["code"] == "UPLOAD_INCOMPLETE"
            for error in invalid.json()["validation_report"]["errors"]
        )

        _upload_all(client, invalid.json(), payloads)
        valid = client.post(
            f"/api/v1/upload-sessions/{session['session_id']}/validate"
        )
        assert valid.status_code == 200
        assert valid.json()["state"] == "READY_TO_CREATE"


def test_upload_rejects_size_mismatch_and_discard_removes_session(
    tmp_path: Path,
) -> None:
    app = create_app(settings=_settings(tmp_path))
    payloads = _valid_files()

    with TestClient(app) as client:
        session = _create_session(client, payloads)
        file = session["files"][0]
        mismatch = client.put(
            file["upload_url"],
            content=b"short",
            headers={"Content-Type": _CONTENT_TYPE},
        )
        assert mismatch.status_code == 422
        assert mismatch.json()["error"]["code"] == "INVALID_REQUEST"

        discarded = client.delete(
            f"/api/v1/upload-sessions/{session['session_id']}"
        )
        assert discarded.status_code == 204
        missing = client.get(
            f"/api/v1/upload-sessions/{session['session_id']}"
        )
        assert missing.status_code == 404
        assert (
            missing.json()["error"]["code"]
            == "UPLOAD_SESSION_NOT_FOUND"
        )


def test_upload_manifest_rejects_unsafe_or_incorrect_folder_names(
    tmp_path: Path,
) -> None:
    app = create_app(settings=_settings(tmp_path))

    with TestClient(app) as client:
        unsafe = client.post(
            "/api/v1/upload-sessions",
            json={
                "files": [
                    {
                        "relative_path": "../scope.docx",
                        "size_bytes": 1,
                        "content_type": _CONTENT_TYPE,
                    }
                ]
            },
        )
        incorrect_folder = client.post(
            "/api/v1/upload-sessions",
            json={
                "files": [
                    {
                        "relative_path": "audit-project/AWp/scope.docx",
                        "size_bytes": 1,
                        "content_type": _CONTENT_TYPE,
                    }
                ]
            },
        )

    assert unsafe.status_code == 422
    assert incorrect_folder.status_code == 422
    assert unsafe.json()["error"]["code"] == "INVALID_REQUEST"
    assert incorrect_folder.json()["error"]["code"] == "INVALID_REQUEST"


def test_upload_manifest_ignores_unsupported_files(tmp_path: Path) -> None:
    app = create_app(settings=_settings(tmp_path))
    supported = _docx_bytes("Approved audit scope")

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/upload-sessions",
            json={
                "files": [
                    {
                        "relative_path": "audit-project/AWP/scope.docx",
                        "size_bytes": len(supported),
                        "content_type": _CONTENT_TYPE,
                    },
                    {
                        "relative_path": "audit-project/.DS_Store",
                        "size_bytes": 24,
                        "content_type": "application/octet-stream",
                    },
                    {
                        "relative_path": "audit-project/APM/notes.txt",
                        "size_bytes": 12,
                        "content_type": "text/plain",
                    },
                    {
                        "relative_path": "audit-project/AWP/~$draft.docx",
                        "size_bytes": 162,
                        "content_type": _CONTENT_TYPE,
                    },
                    {
                        "relative_path": "audit-project/AWP/.~draft.docx",
                        "size_bytes": 162,
                        "content_type": _CONTENT_TYPE,
                    },
                ]
            },
        )

    assert response.status_code == 201, response.text
    assert [item["relative_path"] for item in response.json()["files"]] == [
        "audit-project/AWP/scope.docx"
    ]
